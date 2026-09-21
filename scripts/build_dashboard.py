from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "Excel" / "Procurement KPI Analysis Dataset.xlsx"
DATA_SHEET = "Procurement KPI Analysis Datase"

NAVY = "17324D"
BLUE = "277DA1"
TEAL = "43AA8B"
GOLD = "F4A261"
RED = "E76F51"
INK = "203040"
MUTED = "667788"
PALE_BLUE = "EAF3F8"
PALE_TEAL = "E7F4EF"
PALE_GOLD = "FFF1D9"
PALE_RED = "FBE7E3"
WHITE = "FFFFFF"
LIGHT_BORDER = "D6E0E8"


def filtered_sum(field_col, value_col=None):
    criteria = (
        f"('{DATA_SHEET}'!$B$2:$B$778=$B$6)+($B$6=\"All Suppliers\")",
        f"('{DATA_SHEET}'!$E$2:$E$778=$D$6)+($D$6=\"All Categories\")",
        f"('{DATA_SHEET}'!$F$2:$F$778=$F$6)+($F$6=\"All Statuses\")",
        f"('{DATA_SHEET}'!$K$2:$K$778=$H$6)+($H$6=\"All Compliance\")",
    )
    masks = ")*(".join(criteria)
    value_range = value_col or field_col
    return f"=SUMPRODUCT(({masks})*('{DATA_SHEET}'!{value_range}2:{value_range}778))"


def filtered_count():
    criteria = (
        f"('{DATA_SHEET}'!$B$2:$B$778=$B$6)+($B$6=\"All Suppliers\")",
        f"('{DATA_SHEET}'!$E$2:$E$778=$D$6)+($D$6=\"All Categories\")",
        f"('{DATA_SHEET}'!$F$2:$F$778=$F$6)+($F$6=\"All Statuses\")",
        f"('{DATA_SHEET}'!$K$2:$K$778=$H$6)+($H$6=\"All Compliance\")",
    )
    masks = ")*(".join(criteria)
    return f"=SUMPRODUCT(({masks})*1)"


def write_card(ws, cell, title, formula, fill, number_format):
    col = ws[cell].column
    row = ws[cell].row
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + 1)
    title_cell = ws.cell(row, col, title)
    value_cell = ws.cell(row + 1, col, formula)
    title_cell.font = Font(name="Aptos", size=10, bold=True, color=MUTED)
    title_cell.fill = PatternFill("solid", fgColor=fill)
    value_cell.font = Font(name="Aptos Display", size=20, bold=True, color=NAVY)
    value_cell.fill = PatternFill("solid", fgColor=fill)
    value_cell.number_format = number_format
    for r in range(row, row + 3):
        for c in range(col, col + 2):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)
            ws.cell(r, c).border = Border(
                left=Side(style="thin", color=LIGHT_BORDER),
                right=Side(style="thin", color=LIGHT_BORDER),
                top=Side(style="thin", color=LIGHT_BORDER),
                bottom=Side(style="thin", color=LIGHT_BORDER),
            )
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    value_cell.alignment = Alignment(horizontal="left", vertical="center")


def build_dashboard():
    wb = load_workbook(WORKBOOK_PATH)
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0)
    wb.active = wb.index(ws)
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet is ws
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A9"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.tabColor = BLUE

    widths = {"A": 15, "B": 15, "C": 3, "D": 15, "E": 15, "F": 3, "G": 15, "H": 15, "I": 3, "J": 15, "K": 15, "L": 3, "M": 15, "N": 15}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in range(1, 43):
        ws.row_dimensions[row].height = 20
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 26

    ws.merge_cells("A1:N1")
    ws["A1"] = "PROCUREMENT PERFORMANCE DASHBOARD"
    ws["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A2:N2")
    ws["A2"] = "Spend, supplier performance, quality and order fulfilment | Potential savings are indicative"
    ws["A2"].font = Font(name="Aptos", size=10, italic=True, color=WHITE)
    ws["A2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A4:N4")
    ws["A4"] = "FILTERS"
    ws["A4"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    ws["A4"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
    filter_labels = [("A5", "Supplier"), ("C5", "Category"), ("E5", "Order status"), ("G5", "Compliance")]
    for label_cell, label in filter_labels:
        ws[label_cell] = label
        ws[label_cell].font = Font(name="Aptos", size=9, bold=True, color=MUTED)
    filter_cells = {"B6": "All Suppliers", "D6": "All Categories", "F6": "All Statuses", "H6": "All Compliance"}
    for cell, value in filter_cells.items():
        ws[cell] = value
        ws[cell].font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        ws[cell].fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws[cell].alignment = Alignment(horizontal="left", vertical="center")
        ws[cell].border = Border(bottom=Side(style="medium", color=BLUE))
    ws.merge_cells("J5:N6")
    ws["J5"] = "Use the four selectors to focus every KPI and chart."
    ws["J5"].font = Font(name="Aptos", size=9, color=MUTED)
    ws["J5"].alignment = Alignment(wrap_text=True, vertical="center")

    validations = [
        ("B6", '"All Suppliers,Alpha_Inc,Beta_Supplies,Delta_Logistics,Epsilon_Group,Gamma_Co"'),
        ("D6", '"All Categories,Electronics,MRO,Office Supplies,Packaging,Raw Materials"'),
        ("F6", '"All Statuses,Cancelled,Delivered,Partially Delivered,Pending"'),
        ("H6", '"All Compliance,Yes,No"'),
    ]
    for cell, formula in validations:
        dv = DataValidation(type="list", formula1=formula, allow_blank=False)
        dv.error = "Choose a value from the list."
        dv.errorTitle = "Invalid filter"
        ws.add_data_validation(dv)
        dv.add(ws[cell])

    write_card(ws, "A9", "PURCHASE ORDERS", filtered_count(), PALE_BLUE, "#,##0")
    write_card(ws, "D9", "NEGOTIATED SPEND", filtered_sum("$M", "$M"), PALE_TEAL, "$#,##0")
    write_card(ws, "G9", "POTENTIAL SAVINGS", filtered_sum("$N", "$N"), PALE_GOLD, "$#,##0")
    write_card(ws, "J9", "SAVINGS RATE", "=IFERROR(G10/(G10+D10),0)", PALE_GOLD, "0.0%")
    write_card(ws, "M9", "COMPLIANCE RATE", "=IFERROR(SUMPRODUCT((('" + DATA_SHEET + "'!$B$2:$B$778=$B$6)+($B$6=\"All Suppliers\"))*(('" + DATA_SHEET + "'!$E$2:$E$778=$D$6)+($D$6=\"All Categories\"))*(('" + DATA_SHEET + "'!$F$2:$F$778=$F$6)+($F$6=\"All Statuses\"))*(('" + DATA_SHEET + "'!$K$2:$K$778=$H$6)+($H$6=\"All Compliance\"))*('" + DATA_SHEET + "'!$K$2:$K$778=\"Yes\"))/A10,0)", PALE_TEAL, "0.0%")

    ws.merge_cells("A14:N14")
    ws["A14"] = "PERFORMANCE VIEW"
    ws["A14"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    ws["A14"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A14"].alignment = Alignment(horizontal="left", vertical="center")

    # Helper tables drive the charts and remain hidden from the presentation view.
    helper = {
        "P": ("Supplier", ["Alpha_Inc", "Beta_Supplies", "Delta_Logistics", "Epsilon_Group", "Gamma_Co"]),
        "T": ("Category", ["Electronics", "MRO", "Office Supplies", "Packaging", "Raw Materials"]),
        "X": ("Status", ["Delivered", "Partially Delivered", "Pending", "Cancelled"]),
    }
    for col, (header, labels) in helper.items():
        ws[f"{col}30"] = header
        for index, label in enumerate(labels, 31):
            ws[f"{col}{index}"] = label
    for row in range(30, 36):
        ws[f"Q{row}"] = "Negotiated Spend"
        ws[f"R{row}"] = "Potential Savings"
    for row in range(30, 36):
        ws[f"U{row}"] = "Defect Rate"
        ws[f"V{row}"] = "Compliance Rate"
    for row in range(30, 35):
        ws[f"Y{row}"] = "PO Count"

    for row in range(31, 36):
        label = f"$P{row}"
        mask = f"((('{DATA_SHEET}'!$B$2:$B$778={label})+($B$6=\"All Suppliers\"))*((('{DATA_SHEET}'!$E$2:$E$778=$D$6)+($D$6=\"All Categories\")))*((('{DATA_SHEET}'!$F$2:$F$778=$F$6)+($F$6=\"All Statuses\")))*((('{DATA_SHEET}'!$K$2:$K$778=$H$6)+($H$6=\"All Compliance\"))))"
        ws[f"Q{row}"] = f"=SUMPRODUCT({mask}*'{DATA_SHEET}'!$M$2:$M$778)"
        ws[f"R{row}"] = f"=SUMPRODUCT({mask}*'{DATA_SHEET}'!$N$2:$N$778)"
    for row in range(31, 36):
        label = f"$T{row}"
        mask = f"((('{DATA_SHEET}'!$E$2:$E$778={label})+($D$6=\"All Categories\"))*((('{DATA_SHEET}'!$B$2:$B$778=$B$6)+($B$6=\"All Suppliers\")))*((('{DATA_SHEET}'!$F$2:$F$778=$F$6)+($F$6=\"All Statuses\")))*((('{DATA_SHEET}'!$K$2:$K$778=$H$6)+($H$6=\"All Compliance\"))))"
        total_qty = f"SUMPRODUCT({mask}*'{DATA_SHEET}'!$G$2:$G$778)"
        ws[f"U{row}"] = f"=IFERROR(SUMPRODUCT({mask}*'{DATA_SHEET}'!$J$2:$J$778)/({total_qty}),0)"
        ws[f"V{row}"] = f"=IFERROR(SUMPRODUCT({mask}*('{DATA_SHEET}'!$K$2:$K$778=\"Yes\"))/SUMPRODUCT({mask}),0)"
    for row in range(31, 35):
        label = f"$X{row}"
        mask = f"((('{DATA_SHEET}'!$F$2:$F$778={label})+($F$6=\"All Statuses\"))*((('{DATA_SHEET}'!$B$2:$B$778=$B$6)+($B$6=\"All Suppliers\")))*((('{DATA_SHEET}'!$E$2:$E$778=$D$6)+($D$6=\"All Categories\")))*((('{DATA_SHEET}'!$K$2:$K$778=$H$6)+($H$6=\"All Compliance\"))))"
        ws[f"Y{row}"] = f"=SUMPRODUCT({mask}*1)"
    for col in ["P", "Q", "R", "T", "U", "V", "X", "Y"]:
        ws.column_dimensions[col].hidden = True

    spend_chart = BarChart()
    spend_chart.type = "bar"
    spend_chart.style = 10
    spend_chart.title = "Supplier spend and indicative savings"
    spend_chart.y_axis.title = "Supplier"
    spend_chart.x_axis.title = "Amount"
    spend_chart.height = 7.2
    spend_chart.width = 14.5
    spend_chart.visible_cells_only = False
    spend_chart.add_data(Reference(ws, min_col=17, max_col=18, min_row=30, max_row=35), titles_from_data=True)
    spend_chart.set_categories(Reference(ws, min_col=16, min_row=31, max_row=35))
    spend_chart.varyColors = False
    spend_chart.series[0].graphicalProperties.solidFill = BLUE
    spend_chart.series[1].graphicalProperties.solidFill = GOLD
    ws.add_chart(spend_chart, "A16")

    quality_chart = BarChart()
    quality_chart.type = "col"
    quality_chart.style = 11
    quality_chart.title = "Category quality and compliance"
    quality_chart.y_axis.title = "Rate"
    quality_chart.height = 7.2
    quality_chart.width = 14.5
    quality_chart.visible_cells_only = False
    quality_chart.add_data(Reference(ws, min_col=21, max_col=22, min_row=30, max_row=35), titles_from_data=True)
    quality_chart.set_categories(Reference(ws, min_col=20, min_row=31, max_row=35))
    quality_chart.y_axis.numFmt = "0%"
    quality_chart.series[0].graphicalProperties.solidFill = RED
    quality_chart.series[1].graphicalProperties.solidFill = TEAL
    ws.add_chart(quality_chart, "H16")

    status_chart = DoughnutChart()
    status_chart.title = "Order status mix"
    status_chart.holeSize = 58
    status_chart.height = 7.2
    status_chart.width = 14.5
    status_chart.visible_cells_only = False
    status_chart.add_data(Reference(ws, min_col=25, min_row=30, max_row=34), titles_from_data=True)
    status_chart.set_categories(Reference(ws, min_col=24, min_row=31, max_row=34))
    status_chart.dLbls = DataLabelList()
    status_chart.dLbls.showPercent = True
    status_chart.dLbls.showLeaderLines = True
    for index, color in enumerate([TEAL, GOLD, BLUE, RED]):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color
        status_chart.series[0].data_points.append(point)
    ws.add_chart(status_chart, "A31")

    # A compact narrative panel keeps the dashboard decision-oriented.
    ws.merge_cells("H31:N31")
    ws["H31"] = "READOUT"
    ws["H31"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    ws["H31"].fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells("H32:N39")
    ws["H32"] = (
        "Review the supplier chart for spend concentration and indicative savings. "
        "Use the quality view to compare defect and compliance rates by category. "
        "The status mix highlights fulfilment exposure. Filter results before making supplier or category decisions."
    )
    ws["H32"].font = Font(name="Aptos", size=10, color=INK)
    ws["H32"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["H32"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    for row in range(32, 40):
        for col in range(8, 15):
            ws.cell(row, col).border = Border(
                left=Side(style="thin", color=LIGHT_BORDER),
                right=Side(style="thin", color=LIGHT_BORDER),
                top=Side(style="thin", color=LIGHT_BORDER),
                bottom=Side(style="thin", color=LIGHT_BORDER),
            )

    ws.auto_filter.ref = "A1:N39"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(WORKBOOK_PATH)


if __name__ == "__main__":
    build_dashboard()